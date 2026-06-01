from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK


SOURCE_SECTIONS = [
    (
        "Dataset đầu vào",
        [
            r"C:\KLTN\KLTN\data\final\vismishds_phase1_final.csv",
            r"C:\KLTN\KLTN\model\base\temp.csv",
        ],
    ),
    (
        "Setup A/B - baseline huấn luyện trên real data",
        [
            r"C:\KLTN\KLTN\setup_results\setup_a_results\setup_a_results.csv",
            r"C:\KLTN\KLTN\setup_results\setup_a_results\setup_a_real_test.csv",
            r"C:\KLTN\KLTN\setup_results\setup_b_results\setup_b_results.csv",
            r"C:\KLTN\KLTN\setup_results\setup_b_results\setup_b_per_seed_results.csv",
            r"C:\KLTN\KLTN\notebooks\script_py\setup_a_trtr_phobert.py",
            r"C:\KLTN\KLTN\notebooks\script_py\setup_b_trtr_imbalance_phobert.py",
        ],
    ),
    (
        "Setup C/D - TSTR, train on synthetic và test on real",
        [
            r"C:\KLTN\KLTN\setup_results\setup_c_results\setup_c_results.csv",
            r"C:\KLTN\KLTN\setup_results\setup_c_results\setup_c_per_seed_results.csv",
            r"C:\KLTN\KLTN\setup_results\setup_c_results\setup_c_leakage_report.csv",
            r"C:\KLTN\KLTN\setup_results\setup_c_results\setup_c_synthetic_train_matched.csv",
            r"C:\KLTN\KLTN\setup_results\setup_d_results\setup_d_results.csv",
            r"C:\KLTN\KLTN\setup_results\setup_d_results\setup_d_per_seed_results.csv",
            r"C:\KLTN\KLTN\setup_results\setup_d_results\setup_d_leakage_report.csv",
            r"C:\KLTN\KLTN\setup_results\setup_d_results\setup_d_synthetic_train_balanced.csv",
            r"C:\KLTN\KLTN\notebooks\script_py\setup_c_tstr_matched_phobert.py",
            r"C:\KLTN\KLTN\notebooks\script_py\setup_d_tstr_balanced_phobert.py",
        ],
    ),
    (
        "Setup E - Real Train cộng synthetic Label 1",
        [
            r"C:\KLTN\KLTN\setup_results\setup_e_results\setup_e_results.csv",
            r"C:\KLTN\KLTN\setup_results\setup_e_results\setup_e_per_seed_results.csv",
            r"C:\KLTN\KLTN\setup_results\setup_e_results\setup_e_variant_metadata.csv",
            r"C:\KLTN\KLTN\setup_results\setup_e_results\setup_e_leakage_report.csv",
            r"C:\KLTN\KLTN\setup_results\setup_e_results\setup_e_train_E1_500.csv",
            r"C:\KLTN\KLTN\setup_results\setup_e_results\setup_e_train_E2_1000.csv",
            r"C:\KLTN\KLTN\setup_results\setup_e_results\setup_e_train_E3_2000.csv",
            r"C:\KLTN\KLTN\setup_results\setup_e_results\setup_e_train_E4_all.csv",
            r"C:\KLTN\KLTN\notebooks\script_py\setup_e_real_plus_synthetic_label1_phobert.py",
        ],
    ),
    (
        "Setup F - augmentation nguồn Label 0",
        [
            r"C:\KLTN\KLTN\setup_results\setup_f_results\setup_f_partial_results.csv",
            r"C:\KLTN\KLTN\setup_results\setup_f_results\setup_f_partial_per_seed_results.csv",
            r"C:\KLTN\KLTN\setup_results\setup_f_results\setup_f_variant_metadata.csv",
            r"C:\KLTN\KLTN\setup_results\setup_f_results\setup_f_leakage_report.csv",
            r"C:\KLTN\KLTN\setup_results\setup_f_results\setup_f3_results.csv",
            r"C:\KLTN\KLTN\setup_results\setup_f_results\setup_f3_per_seed_results.csv",
            r"C:\KLTN\KLTN\setup_results\setup_f_results\setup_f3_variant_metadata.csv",
            r"C:\KLTN\KLTN\setup_results\setup_f_results\setup_f3_leakage_report.csv",
            r"C:\KLTN\KLTN\setup_results\setup_f_results\setup_f3_train.csv",
            r"C:\KLTN\KLTN\notebooks\script_py\setup_f_label0_source_augmentation_phobert.py",
            r"C:\KLTN\KLTN\notebooks\script_py\setup_f3_kaggle_label0_source_augmentation_phobert.py",
        ],
    ),
    (
        "Setup G - external challenge và kiểm tra ổn định miền Label 0",
        [
            r"C:\KLTN\KLTN\setup_results\setup_g_results\setup_g_results.csv",
            r"C:\KLTN\KLTN\setup_results\setup_g_results\setup_g_per_seed_results.csv",
            r"C:\KLTN\KLTN\setup_results\setup_g_results\setup_g_subset_results.csv",
            r"C:\KLTN\KLTN\setup_results\setup_g_results\setup_g_subset_summary.csv",
            r"C:\KLTN\KLTN\setup_results\setup_g_results\setup_g_variant_metadata.csv",
            r"C:\KLTN\KLTN\setup_results\setup_g_results\setup_g_synthetic_leakage_report.csv",
            r"C:\KLTN\KLTN\setup_results\setup_g_results\setup_g_challenge_split_metadata.csv",
            r"C:\KLTN\KLTN\setup_results\setup_g_results\setup_g_challenge_train.csv",
            r"C:\KLTN\KLTN\setup_results\setup_g_results\setup_g_challenge_val.csv",
            r"C:\KLTN\KLTN\setup_results\setup_g_results\setup_g_challenge_test.csv",
            r"C:\KLTN\KLTN\notebooks\script_py\setup_g_external_challenge_phobert.py",
        ],
    ),
    (
        "Tài liệu tổng hợp trong repo",
        [
            r"C:\KLTN\KLTN\docs\tstr_dataset_evaluation_plan.md",
            r"C:\KLTN\KLTN\docs\ViSmishDS_Data_Augmentation.md",
        ],
    ),
]


def remove_existing_source_section(doc: Document) -> None:
    start_idx = None
    for i, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == "Nguồn số liệu và artefact tái lập":
            start_idx = i
            break
    if start_idx is None:
        return
    for paragraph in doc.paragraphs[start_idx:]:
        element = paragraph._element
        element.getparent().remove(element)


def add_sources(doc: Document) -> None:
    if doc.paragraphs:
        doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Nguồn số liệu và artefact tái lập", level=1)
    doc.add_paragraph(
        "Các số liệu trong phần báo cáo TSTR/augmentation được tổng hợp từ các file kết quả, "
        "metadata, leakage report và script thí nghiệm dưới đây. Đường dẫn là đường dẫn cục bộ "
        "trong repo tại thời điểm viết báo cáo."
    )
    for title, paths in SOURCE_SECTIONS:
        doc.add_heading(title, level=2)
        for path in paths:
            doc.add_paragraph(path, style="List Bullet")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx_path", type=Path)
    parser.add_argument("--backup", action="store_true")
    args = parser.parse_args()

    doc = Document(args.docx_path)
    if args.backup:
        backup_path = args.docx_path.with_name(args.docx_path.stem + "_before_sources.docx")
        doc.save(backup_path)

    remove_existing_source_section(doc)
    add_sources(doc)
    doc.save(args.docx_path)


if __name__ == "__main__":
    main()
