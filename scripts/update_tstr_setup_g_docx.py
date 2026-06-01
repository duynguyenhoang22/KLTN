from pathlib import Path

import pandas as pd
from docx import Document
from docx.shared import Inches


DOCX_PATH = Path(r"C:\Users\duyng\Desktop\TSTR.docx")
OUTPUT_PATH = Path(r"C:\KLTN\KLTN\TSTR_setup_g.docx")
RESULT_DIR = Path(r"C:\KLTN\KLTN\setup_results\setup_g_results")


def set_paragraph_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def paragraph_after(doc, target_el, text="", style=None, bold=False, italic=False):
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    target_el.addnext(paragraph._p)
    return paragraph._p, paragraph


def table_after(doc, target_el, rows, cols):
    table = doc.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    target_el.addnext(table._tbl)
    return table._tbl, table


def picture_after(doc, target_el, image_path, caption):
    target_el, caption_para = paragraph_after(doc, target_el, caption)
    if caption_para.runs:
        caption_para.runs[0].bold = True
    target_el, image_para = paragraph_after(doc, target_el, "")
    image_para.alignment = 1
    image_para.add_run().add_picture(str(image_path), width=Inches(6.4))
    return target_el


def metric_text(df, variant, metric):
    row = df[(df["variant"] == variant) & (df["metric"] == metric)].iloc[0]
    return f'{row["mean"]:.4f} +/- {row["std"]:.4f}'


def fill_table(table, headers, data_rows):
    for idx, header in enumerate(headers):
        table.cell(0, idx).text = header
    for r_idx, row in enumerate(data_rows, start=1):
        for c_idx, value in enumerate(row):
            table.cell(r_idx, c_idx).text = str(value)


def main():
    doc = Document(DOCX_PATH)
    results = pd.read_csv(RESULT_DIR / "setup_g_results.csv")
    variants = pd.read_csv(RESULT_DIR / "setup_g_variant_metadata.csv")

    paragraphs = list(doc.paragraphs)
    setup_anchor = next(p for p in paragraphs if p.text.strip().startswith("Trong Setup E, Real Train"))
    result_anchor = next(p for p in paragraphs if p.text.strip().startswith("Tóm lại, không có một biến thể E"))
    cm_anchor = next(p for p in paragraphs if p.text.strip().startswith("Vì vậy, phân tích confusion matrix"))

    target_el, _ = paragraph_after(
        doc,
        setup_anchor._p,
        "Setup G - External Challenge Evaluation",
        style="List Paragraph",
    )
    setup_paragraphs = [
        "Setup G được bổ sung để kiểm tra trực tiếp giới hạn của kết luận từ Setup E: mô hình E4 hoạt động tốt trên Real Test hiện có, nhưng chưa chắc ổn định khi tập test được mở rộng bằng các mẫu external thuộc miền ViLexNorm. Vì vậy, Setup G tái tạo lại công thức E4 champion làm baseline G0, sau đó thêm các nguồn external Label 0 và synthetic Label 0 để đánh giá khả năng giữ ranh giới phân loại khi miền Label 0 được mở rộng.",
        "Challenge Test của Setup G gồm Real Test đã khóa từ Setup A/E kết hợp với external_test. Cụ thể, challenge_test có 537 mẫu, gồm 386 mẫu real và 151 mẫu external; phân phối nhãn là 500 Label 0 và 37 Label 1. External chỉ chứa Label 0, nên đây là phép kiểm tra rất tập trung vào false positive: mô hình có dự đoán nhầm các tin nhắn hợp lệ ngoài miền hiện có thành smishing hay không.",
        "Các biến thể Setup G gồm G0_E4_champion, G1_external_all, G2_external_curated và G3_external_plus_synthetic_label0. Tất cả đều dùng PhoBERT base, cùng 3 seed 42, 123 và 2025. G0 giữ validation theo E4 real validation để tái hiện baseline cũ; G1-G3 dùng challenge validation có thêm external validation để chọn checkpoint phù hợp hơn với miền đánh giá mở rộng.",
    ]
    for text in setup_paragraphs:
        target_el, _ = paragraph_after(doc, target_el, text)

    variant_rows = []
    for _, row in variants.iterrows():
        variant_rows.append(
            [
                row["variant"],
                int(row["train_total"]),
                int(row["train_label0"]),
                int(row["train_label1"]),
                int(row["train_external_real"]),
                int(row["train_external_curated"]),
                int(row["train_synthetic"]),
                row["validation_frame"],
            ]
        )
    target_el, table = table_after(doc, target_el, rows=len(variant_rows) + 1, cols=8)
    fill_table(
        table,
        [
            "Variant",
            "Train total",
            "Label 0",
            "Label 1",
            "External real",
            "External curated",
            "Synthetic",
            "Validation",
        ],
        variant_rows,
    )

    target_el, _ = paragraph_after(
        doc,
        result_anchor._p,
        "Kết Quả Setup G",
        style="List Paragraph",
    )
    result_paragraphs = [
        "Khi mở rộng Real Test bằng external_test, G0_E4_champion giảm rõ rệt so với kết quả E4 trên Real Test thuần. G0 chỉ đạt Macro-F1 = 0.7902 +/- 0.0295 và F1 Label 1 = 0.6237 +/- 0.0503 trên challenge_all, dù Recall Label 1 vẫn cao ở mức 0.9189 +/- 0.0220. Nguyên nhân chính là Precision Label 1 giảm xuống 0.4745 +/- 0.0587, cho thấy mô hình E4 dự đoán quá nhiều mẫu Label 0 external thành Label 1.",
        "Các biến thể có bổ sung external Label 0 cải thiện rất mạnh. G1_external_all đạt F1 Label 1 = 0.8924 +/- 0.0149 và Macro-F1 = 0.9422 +/- 0.0081. G2_external_curated là biến thể tốt nhất theo Macro-F1 và F1 Label 1, lần lượt đạt 0.9470 +/- 0.0073 và 0.9014 +/- 0.0133. G3_external_plus_synthetic_label0 đạt kết quả gần G1/G2, nhưng không vượt G2, cho thấy thêm synthetic Label 0 chưa tạo cải thiện rõ ràng trong thiết lập này.",
        "Kết quả Setup G làm rõ rằng E4 là lựa chọn tốt khi đánh giá trên Real Test hiện có, nhưng chưa đủ ổn định khi xuất hiện miền Label 0 external. Bổ sung external_curated Label 0 giúp mô hình học ranh giới hợp lệ/smishing tốt hơn mà không cần đưa toàn bộ external_real vào train; vì vậy G2_external_curated là lựa chọn cân bằng nhất trong nhóm Setup G.",
    ]
    for text in result_paragraphs:
        target_el, _ = paragraph_after(doc, target_el, text)

    metric_rows = []
    for variant in [
        "G0_E4_champion",
        "G1_external_all",
        "G2_external_curated",
        "G3_external_plus_synthetic_label0",
    ]:
        metric_rows.append(
            [
                variant,
                metric_text(results, variant, "macro_f1"),
                metric_text(results, variant, "f1_label1"),
                metric_text(results, variant, "recall_label1"),
                metric_text(results, variant, "precision_label1"),
                metric_text(results, variant, "auprc"),
                metric_text(results, variant, "fpr_label0"),
            ]
        )
    target_el, table = table_after(doc, target_el, rows=len(metric_rows) + 1, cols=7)
    fill_table(
        table,
        [
            "Variant",
            "Macro-F1",
            "F1 L1",
            "Recall L1",
            "Precision L1",
            "AUPRC",
            "FPR L0",
        ],
        metric_rows,
    )

    target_el, _ = paragraph_after(
        doc,
        cm_anchor._p,
        "Confusion Matrix Setup G",
        style="List Paragraph",
    )
    cm_text = (
        "Các confusion matrix của Setup G cho thấy nguyên nhân suy giảm của G0 nằm ở false positive trên nhóm external Label 0. "
        "Khi external Label 0 được đưa vào train, đặc biệt ở G2_external_curated, số mẫu external bị dự đoán nhầm thành Label 1 giảm mạnh, "
        "trong khi khả năng nhận diện Label 1 trên phần real của challenge test vẫn được giữ ở mức cao."
    )
    target_el, _ = paragraph_after(doc, target_el, cm_text)

    for variant, label in [
        ("G0_E4_champion", "G0 - E4 champion"),
        ("G1_external_all", "G1 - external all"),
        ("G2_external_curated", "G2 - external curated"),
        ("G3_external_plus_synthetic_label0", "G3 - external all + synthetic Label 0"),
    ]:
        image = RESULT_DIR / f"setup_g_confusion_matrix_{variant}_challenge_all.png"
        target_el = picture_after(doc, target_el, image, f"Confusion matrix Setup G trên challenge_all: {label}")

    target_el, _ = paragraph_after(
        doc,
        target_el,
        "Bên cạnh challenge_all, các ma trận trên external_all giúp đọc riêng mức độ false positive của từng biến thể vì external_test chỉ gồm Label 0. G0 có tỷ lệ dự đoán nhầm external thành Label 1 cao nhất, còn G1-G3 kiểm soát lỗi này tốt hơn nhiều.",
    )
    for variant, label in [
        ("G0_E4_champion", "G0 - E4 champion"),
        ("G1_external_all", "G1 - external all"),
        ("G2_external_curated", "G2 - external curated"),
        ("G3_external_plus_synthetic_label0", "G3 - external all + synthetic Label 0"),
    ]:
        image = RESULT_DIR / f"setup_g_confusion_matrix_{variant}_external_all.png"
        target_el = picture_after(doc, target_el, image, f"Confusion matrix Setup G trên external_all: {label}")

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Cuối cùng, Setup F chưa được hoàn tất."):
            set_paragraph_text(
                paragraph,
                "Sau khi bổ sung Setup G, hạn chế chính còn lại là tập Real Test và challenge_test vẫn có số lượng Label 1 nhỏ, chỉ 37 mẫu. External trong Setup G hiện chỉ đóng vai trò Label 0, nên thí nghiệm này chủ yếu kiểm tra false positive và độ ổn định của ranh giới Label 0, chưa đánh giá được các biến thể smishing external thuộc Label 1.",
            )
        elif text.startswith("Dữ liệu tạo sinh không nên được xem là alternative data"):
            set_paragraph_text(
                paragraph,
                "Dữ liệu tạo sinh không nên được xem là alternative data thay thế dữ liệu thật, mà nên được xem là data augmentation có kiểm soát. Setup E cho thấy synthetic Label 1 có thể hỗ trợ lớp smishing khi vẫn giữ Real Train làm điểm neo. Setup G bổ sung thêm bằng chứng rằng khi miền Label 0 mở rộng, external_curated Label 0 giúp ổn định ranh giới phân loại tốt hơn, trong đó G2_external_curated là biến thể cân bằng nhất trên challenge_test.",
            )
        elif text == "Setup F sẽ tiếp tục kiểm tra các variant:":
            set_paragraph_text(
                paragraph,
                "Setup G đã hiện thực nhóm thí nghiệm external challenge với các variant:",
            )
        elif text == "F1: Synthetic Label 0 đối chứng":
            set_paragraph_text(paragraph, "G0: E4 champion, không thêm external vào train")
        elif text == "F2a: external_real Label 0":
            set_paragraph_text(paragraph, "G1: external_real + external_curated Label 0")
        elif text == "F2b: external_curated Label 0":
            set_paragraph_text(paragraph, "G2: external_curated Label 0")
        elif text == "F2c: external_real + external_curated Label 0":
            set_paragraph_text(paragraph, "G3: external_real + external_curated Label 0 + synthetic Label 0")
        elif text == "F3: external all + synthetic Label 0":
            set_paragraph_text(paragraph, "")
        elif text.startswith("nhằm đánh giá liệu bổ sung cả hai lớp"):
            set_paragraph_text(
                paragraph,
                "Kết quả cho thấy external Label 0, đặc biệt external_curated, có giá trị rõ rệt trong việc giảm false positive khi đánh giá trên challenge_test có miền Label 0 mở rộng.",
            )

    doc.save(OUTPUT_PATH)


if __name__ == "__main__":
    main()
